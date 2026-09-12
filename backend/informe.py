"""
El marco de los informes: UNA hoja, el mismo formato para todas las áreas.

QUÉ PROBLEMA RESUELVE
---------------------
Lo que se descargaba antes era una tabla plana de tres columnas —«Sección | Detalle |
Valor»— con el mismo molde para todas las pantallas. Se leía como una foto de la lista
que ya estaba en pantalla: sin totales, sin porcentajes, sin comparación, y con secciones
enteras fuera. Un número suelto —«EXTERNO 18»— no dice nada: ni de cuántos, ni contra qué
mes, ni si eso es mucho o poco para ese bote.

UNA DEFINICIÓN, DOS ARCHIVOS
----------------------------
Cada área describe su informe UNA vez, en un diccionario (ver informes_areas.py), y de
ahí salen el PDF y el libro de Excel. Es lo que garantiza que los dos digan lo mismo: si
salieran de dos códigos distintos, tarde o temprano uno se quedaría viejo y nadie sabría
cuál creer. Y es lo que hace que el informe de Restaurantes se lea igual que el de
Analítica, que es justo lo que se pidió.

El PDF cabe en una hoja a propósito. La regla al armarlo: si algo no cabe, no es que haya
que encogerlo, es que no entra — entra lo que se mira para decidir, y el detalle completo
va en la hoja «Data» del libro de Excel.

POR DÍA O POR MES, EL MISMO FORMATO
-----------------------------------
El informe recibe un rango cualquiera. Lo único que se adapta es el bloque de tendencia,
porque «por semana» no significa nada en un informe de un día:

  · hasta 21 días  -> una barra por DÍA. Si el rango es más corto que una semana, se
    extiende hacia atrás hasta completar siete barras, y las de antes del periodo van en
    tono claro: son contexto, no son el periodo.
  · más de 21 días -> una barra por SEMANA.

La comparación también se adapta: si el rango es un mes de calendario exacto se compara
con el mes anterior de calendario; si no, con la misma cantidad de días justo antes.

EN INGLÉS
---------
Los informes salen SIEMPRE en inglés: se mandan a inversionistas y socios que no leen
español. La pantalla del sistema sigue en español, y los comentarios de este archivo
también — lo que cambia es lo que se imprime. Los nombres propios que vienen de los datos
—Sierpe, Drake, Terra Kitchen, los tours del catálogo— se dejan como están: son nombres,
no texto traducible.

NADA DE DINERO
--------------
No llevan tarifas, precios ni montos, por decisión de operación. Las CORTESÍAS sí van,
pero como CUÁNTAS, que es el número que se quiere seguir.
"""
import calendar
import datetime
import io
import os
import re

from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_RIGHT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import (BaseDocTemplate, Frame, Image, PageTemplate,
                                Paragraph, Spacer, Table, TableStyle)

# ---------------------------------------------------------------------------
# Identidad. Los mismos colores de la pantalla.
# ---------------------------------------------------------------------------
BOSQUE = colors.HexColor("#2C4A3E")
VERDE_MEDIO = colors.HexColor("#5A7A6B")
VERDE_PALIDO = colors.HexColor("#B9C9BF")
ARENA = colors.HexColor("#F7F4EC")
LINEA = colors.HexColor("#DCD7C9")
TINTA = colors.HexColor("#23241F")
GRIS = colors.HexColor("#6E6A5C")
BAJA = colors.HexColor("#9C5B2E")
SUBE = colors.HexColor("#3F7A52")

MARGEN = 13 * mm
ANCHO = A4[0] - 2 * MARGEN
HUECO = 6 * mm
COL = (ANCHO - HUECO) / 2

T = ParagraphStyle("t", fontName="Helvetica", fontSize=7.8, leading=9.6, textColor=TINTA)
T_D = ParagraphStyle("td", parent=T, alignment=TA_RIGHT)
T_B = ParagraphStyle("tb", parent=T, fontName="Helvetica-Bold")
T_BD = ParagraphStyle("tbd", parent=T_B, alignment=TA_RIGHT)
MUTE = ParagraphStyle("m", parent=T, fontSize=6.8, leading=8.2, textColor=GRIS)
H_SEC = ParagraphStyle("hs", fontName="Helvetica-Bold", fontSize=9, leading=10.5,
                       textColor=BOSQUE)

# Los informes se imprimen en inglés (ver el encabezado del módulo).
MESES = ["Jan", "Feb", "Mar", "Apr", "May", "Jun",
         "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
MESES_LARGOS = ["January", "February", "March", "April", "May", "June", "July",
                "August", "September", "October", "November", "December"]
SIN_ASIGNAR = "(unassigned)"
OTROS = "Others"
TOTAL = "TOTAL"
PIE_LIBRO = ("Corcovado Wilderness Lodge · full detail in the Excel workbook "
             "for this same period")
ALCANCE = ("Scope: operational figures from the lodge system (reservations, activities, "
           "transfers and spa). No financial information is included.")

# Hasta aquí se grafica por día; de aquí en adelante, por semana.
TOPE_DIARIO = 21
MINIMO_BARRAS = 7


def _logo():
    ruta = os.path.join(os.path.dirname(__file__), "..", "frontend", "assets", "logo.jpg")
    return ruta if os.path.exists(ruta) else None


def sql_fecha(col):
    """La MISMA expresión que main.sql_fecha: reordena 'DD-MM-YY' a 'YY-MM-DD'."""
    return f"(substr({col},7,2)||'-'||substr({col},4,2)||'-'||substr({col},1,2))"


def yymmdd(iso):
    y, m, d = iso.split("-")
    return f"{y[2:]}-{m}-{d}"


def ddmmyy(iso):
    y, m, d = iso.split("-")
    return f"{d}-{m}-{y[2:]}"


def fecha_larga(iso, con_ano=True):
    d = datetime.date.fromisoformat(iso)
    return f"{d.day} {MESES[d.month - 1]}" + (f" {d.year}" if con_ano else "")


def es_mes_completo(desde, hasta):
    d0, d1 = datetime.date.fromisoformat(desde), datetime.date.fromisoformat(hasta)
    return (d0.day == 1 and d0.year == d1.year and d0.month == d1.month
            and d1.day == calendar.monthrange(d1.year, d1.month)[1])


def titulo_del_periodo(desde, hasta):
    """Cómo se llama este periodo en la cabecera. Un mes entero se dice por su nombre."""
    d0, d1 = datetime.date.fromisoformat(desde), datetime.date.fromisoformat(hasta)
    if d0 == d1:
        return f"{MESES_LARGOS[d0.month - 1]} {d0.day}, {d0.year}"
    if es_mes_completo(desde, hasta):
        return f"{MESES_LARGOS[d0.month - 1]} {d0.year}"
    return f"{fecha_larga(desde, False)} — {fecha_larga(hasta)}"


def periodo_anterior(desde, hasta):
    """Contra qué se compara. Un mes se compara con el mes anterior, no con 30 días."""
    d0, d1 = datetime.date.fromisoformat(desde), datetime.date.fromisoformat(hasta)
    if es_mes_completo(desde, hasta):
        fin = d0 - datetime.timedelta(days=1)
        return fin.replace(day=1).isoformat(), fin.isoformat()
    dias = (d1 - d0).days + 1
    return ((d0 - datetime.timedelta(days=dias)).isoformat(),
            (d0 - datetime.timedelta(days=1)).isoformat())


def plural(n, uno, varios):
    return f"{n} {uno if n == 1 else varios}"


# ---------------------------------------------------------------------------
# Utilidades de datos que usan todas las áreas
# ---------------------------------------------------------------------------

def a_fecha(valor):
    """'05-08-26' -> date. None si no tiene esa forma."""
    try:
        d, m, y = (valor or "").split("-")
        return datetime.date(2000 + int(y), int(m), int(d))
    except (ValueError, AttributeError):
        return None


def total_habitaciones():
    """El mismo total que usa la pantalla de ocupación, para que no se contradigan."""
    try:
        import publicador as _pub
        return len(_pub.cargar_config().get("habitaciones") or []) or 30
    except Exception:
        return 30


def metricas(conn, desde, hasta, habitaciones=None):
    """Las cifras de ocupación y actividad del periodo. Solo lee.

    Vive en el marco y no en un área porque la usan varias: el informe de Analítica, el
    de Restaurantes y el del día. Si cada una la calculara por su lado, dos informes del
    mismo periodo dirían ocupaciones distintas y nadie sabría cuál creer.
    """
    habitaciones = habitaciones or total_habitaciones()
    d0 = datetime.date.fromisoformat(desde)
    d1 = datetime.date.fromisoformat(hasta)
    noches = (d1 - d0).days + 1

    reservas = [dict(r) for r in conn.execute(
        """SELECT conf_no, room_no, arr_date, dep_date, adl, chl, punto_entrada,
                  punto_salida, cortesia, regimen, res_status
           FROM reserva WHERE res_status != 'CANCELADA'""")]

    pax_noche, hab_noche = 0, 0
    for i in range(noches):
        dia = d0 + datetime.timedelta(days=i)
        habs = set()
        for r in reservas:
            llega, sale = a_fecha(r["arr_date"]), a_fecha(r["dep_date"])
            # Se cuenta ocupada la noche que el huésped duerme ahí: desde su llegada
            # hasta el día antes de su salida. Igual que la pantalla de ocupación.
            if llega and llega <= dia and (sale is None or dia < sale):
                habs.add(r["room_no"])
                pax_noche += (r["adl"] or 0) + (r["chl"] or 0)
        hab_noche += len(habs)
    ocupacion = round(100 * hab_noche / (habitaciones * noches)) if noches else 0

    tours = [dict(r) for r in conn.execute(
        """SELECT ta.fecha, ta.tour_codigo, ta.pax, ta.guia_nombre, ta.bote_nombre,
                  r.room_no, r.nombre_principal
           FROM tour_asignado ta
           JOIN reserva r ON r.conf_no = ta.conf_no
           WHERE ta.fecha BETWEEN ? AND ? AND r.res_status != 'CANCELADA'""",
        (desde, hasta))]

    con_tour = {r["conf_no"] for r in conn.execute(
        "SELECT DISTINCT conf_no FROM tour_asignado WHERE fecha BETWEEN ? AND ?",
        (desde, hasta))}
    entraron = [r for r in reservas
                if a_fecha(r["arr_date"]) and d0 <= a_fecha(r["arr_date"]) <= d1]
    salieron = [r for r in reservas
                if a_fecha(r["dep_date"]) and d0 <= a_fecha(r["dep_date"]) <= d1]
    ids_entraron = {r["conf_no"] for r in entraron}

    # Cuántas cortesías, no cuánto valen. Sale de la columna que se llena leyendo la
    # nota de la reserva, así que cuenta lo que Reservaciones escribió.
    cortesias = [r for r in entraron if r.get("cortesia")]

    spa = [dict(r) for r in conn.execute(
        """SELECT c.estado, c.origen, c.minutos, c.terapeuta,
                  COALESCE(s.nombre, c.servicio_codigo) AS tratamiento
           FROM spa_cita c LEFT JOIN spa_servicio s ON s.codigo = c.servicio_codigo
           WHERE c.fecha BETWEEN ? AND ?""", (desde, hasta))]

    return {
        "desde": desde, "hasta": hasta, "noches": noches,
        "pax_noche": pax_noche, "ocupacion": ocupacion,
        "salidas": len(tours), "pax_tours": sum(t["pax"] or 0 for t in tours),
        "reservas_periodo": len(ids_entraron),
        "reservas_con_tour": len(ids_entraron & con_tour),
        "cortesias": len(cortesias),
        "cortesias_pax": sum((r["adl"] or 0) + (r["chl"] or 0) for r in cortesias),
        "entraron": entraron, "salieron": salieron,
        "tours": tours, "spa": spa, "reservas": reservas,
        "spa_activas": len([c for c in spa if c["estado"] in ("HECHA", "CONFIRMADA")]),
    }


def agrupar(filas, clave, campo_pax="pax"):
    salida = {}
    for f in filas:
        k = f.get(clave) or SIN_ASIGNAR
        d = salida.setdefault(k, {"n": 0, "pax": 0})
        d["n"] += 1
        d["pax"] += f.get(campo_pax) or 0
    return salida


def top_y_otros(pares, n, nombre="nombre"):
    """Los n primeros y el resto agrupado.

    Agrupar UNO solo en «Others (1)» no resume nada y esconde un nombre: si sobra uno
    nada más, entra con el suyo.
    """
    ordenados = sorted(pares.items(), key=lambda x: (-x[1]["pax"], -x[1]["n"], str(x[0])))
    if len(ordenados) == n + 1:
        n += 1
    primeros, resto = ordenados[:n], ordenados[n:]
    filas = [{nombre: k, "n": v["n"], "pax": v["pax"]} for k, v in primeros]
    if resto:
        filas.append({nombre: f"{OTROS} ({len(resto)})",
                      "n": sum(v["n"] for _, v in resto),
                      "pax": sum(v["pax"] for _, v in resto)})
    return filas


def serie_de_tendencia(conn, desde, hasta, consulta=None, titulo_dia=None,
                       titulo_semana=None):
    """(valores, etiquetas, dentro, titulo). 'dentro' dice qué barras son del periodo.

    Ver el encabezado del módulo: el cubo se adapta al largo del rango para que el mismo
    bloque sirva en un informe de un día y en uno de un mes. 'consulta' devuelve
    {fecha_iso: valor}; por defecto, el pax en tours.
    """
    d0 = datetime.date.fromisoformat(desde)
    d1 = datetime.date.fromisoformat(hasta)
    dias = (d1 - d0).days + 1

    if dias > TOPE_DIARIO:
        cubos, cursor = [], d0
        while cursor <= d1:
            fin = min(cursor + datetime.timedelta(days=6), d1)
            cubos.append((cursor, fin, f"wk {cursor.strftime('%d/%m')}", True))
            cursor = fin + datetime.timedelta(days=1)
        ini_serie = d0
        titulo = titulo_semana or "Tour pax by week"
    else:
        # Rango corto: se completa hacia atrás para que la barra tenga con qué compararse.
        faltan = max(0, MINIMO_BARRAS - dias)
        ini_serie = d0 - datetime.timedelta(days=faltan)
        cubos, cursor = [], ini_serie
        while cursor <= d1:
            cubos.append((cursor, cursor, cursor.strftime("%d/%m"), cursor >= d0))
            cursor += datetime.timedelta(days=1)
        base = titulo_dia or "Tour pax by day"
        titulo = base if dias > 1 else f"{base} · last {len(cubos)} days"

    if consulta is None:
        def consulta(c, ini, fin):
            return {r["f"]: r["v"] or 0 for r in c.execute(
                """SELECT ta.fecha f, SUM(ta.pax) v FROM tour_asignado ta
                   JOIN reserva r ON r.conf_no = ta.conf_no
                   WHERE ta.fecha BETWEEN ? AND ? AND r.res_status != 'CANCELADA'
                   GROUP BY ta.fecha""", (ini, fin))}

    por_dia = consulta(conn, ini_serie.isoformat(), hasta)
    valores, etiquetas, dentro = [], [], []
    for ini, fin, etiqueta, es_del_periodo in cubos:
        total, cur = 0, ini
        while cur <= fin:
            total += por_dia.get(cur.isoformat(), 0)
            cur += datetime.timedelta(days=1)
        valores.append(total)
        etiquetas.append(etiqueta)
        dentro.append(es_del_periodo)
    return valores, etiquetas, dentro, titulo


# ---------------------------------------------------------------------------
# Piezas del dibujo
# ---------------------------------------------------------------------------

def _p(txt, estilo=T):
    return Paragraph(str(txt), estilo)


def _delta(actual, previo):
    if previo in (None, 0):
        return ("—", GRIS) if not actual else ("new", GRIS)
    cambio = (actual - previo) / previo * 100
    if abs(cambio) < 0.5:
        return "flat", GRIS
    return (f"{'▲' if cambio > 0 else '▼'} {abs(cambio):.0f}%",
            SUBE if cambio > 0 else BAJA)


def seccion(titulo, aclaracion=None, ancho=None):
    t = Table([[_p(titulo.upper(), H_SEC)]], colWidths=[ancho or COL])
    t.setStyle(TableStyle([
        ("LINEBELOW", (0, 0), (-1, 0), 0.8, BOSQUE),
        ("LEFTPADDING", (0, 0), (-1, -1), 0), ("RIGHTPADDING", (0, 0), (-1, -1), 0),
        ("TOPPADDING", (0, 0), (-1, -1), 0), ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
    ]))
    piezas = [t]
    if aclaracion:
        piezas += [Spacer(1, 1.5), _p(aclaracion, MUTE)]
    piezas.append(Spacer(1, 2.5))
    return piezas


def barra(valor, maximo, ancho_mm, color=VERDE_MEDIO):
    llena = 0 if not maximo else max(0.5, ancho_mm * valor / maximo)
    t = Table([["", ""]], colWidths=[llena * mm, max(0.1, ancho_mm - llena) * mm],
              rowHeights=[3.6 * mm])
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (0, 0), color),
        ("BACKGROUND", (1, 0), (1, 0), ARENA),
        ("LEFTPADDING", (0, 0), (-1, -1), 0), ("RIGHTPADDING", (0, 0), (-1, -1), 0),
        ("TOPPADDING", (0, 0), (-1, -1), 0.8), ("BOTTOMPADDING", (0, 0), (-1, -1), 0.8),
    ]))
    return t


def tablita(columnas, filas, anchos, total=None, barra_en=None, tope=None):
    """columnas: [(clave, encabezado, 'i'|'d', formato_excel)]."""
    cab = [_p(f'<font color="white"><b>{c[1]}</b></font>', T_D if c[2] == "d" else T)
           for c in columnas]
    datos = [cab]
    for f in filas:
        fila = []
        for c in columnas:
            if barra_en and c[0] == barra_en:
                fila.append(barra(f.get(c[0]) or 0, tope, anchos[len(fila)] / mm - 1))
                continue
            v = f.get(c[0])
            fila.append(_p("—" if v in (None, "") else v, T_D if c[2] == "d" else T))
        datos.append(fila)
    if total:
        datos.append([_p(f'<b>{total.get(c[0], "")}</b>',
                         T_BD if c[2] == "d" else T_B) for c in columnas])

    t = Table(datos, colWidths=anchos, hAlign="LEFT")
    estilo = [
        ("BACKGROUND", (0, 0), (-1, 0), BOSQUE),
        ("ROWBACKGROUNDS", (0, 1), (-1, -2 if total else -1), [colors.white, ARENA]),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 3.4), ("BOTTOMPADDING", (0, 0), (-1, -1), 3.4),
        ("LEFTPADDING", (0, 0), (-1, -1), 3.5), ("RIGHTPADDING", (0, 0), (-1, -1), 3.5),
    ]
    if total:
        estilo += [("LINEABOVE", (0, -1), (-1, -1), 0.7, BOSQUE),
                   ("BACKGROUND", (0, -1), (-1, -1), colors.white)]
    t.setStyle(TableStyle(estilo))
    return t


def banda_kpi(kpis, neutrales=()):
    """kpis: (etiqueta, valor, previo, sufijo, pie).

    'neutrales' son las cifras donde subir no es bueno ni malo. Las cortesías son el
    caso: pintar de rojo que bajaron da a entender que pasó algo malo, y bajar cortesías
    normalmente es lo contrario. Se muestra el cambio sin opinar.
    """
    p_val = ParagraphStyle("kv", parent=T, fontSize=17, leading=19)
    p_et = ParagraphStyle("ke", parent=T, fontSize=6.2, leading=7.4, textColor=GRIS)
    p_dl = ParagraphStyle("kd", parent=T, fontSize=6.6, leading=8)
    ets, vals, dels = [], [], []
    for etiqueta, valor, previo, sufijo, pie in kpis:
        ets.append(_p(etiqueta.upper(), p_et))
        vals.append(_p(f'<font color="#2C4A3E"><b>{valor}</b></font>'
                       f'<font size="7.5" color="#6E6A5C"> {sufijo}</font>', p_val))
        texto, color = _delta(valor if isinstance(valor, (int, float)) else 0, previo)
        if etiqueta in neutrales:
            color = GRIS
        cola = pie or (f"was {previo}" if previo not in (None, 0) else "")
        dels.append(_p(f'<font color="#{color.hexval()[2:]}"><b>{texto}</b></font>'
                       + (f'<font color="#6E6A5C"> · {cola}</font>' if cola else ""), p_dl))
    t = Table([ets, vals, dels], colWidths=[ANCHO / len(kpis)] * len(kpis),
              rowHeights=[8, 20, 10])
    t.setStyle(TableStyle([
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 0), ("RIGHTPADDING", (0, 0), (-1, -1), 5),
        ("TOPPADDING", (0, 0), (-1, -1), 0), ("BOTTOMPADDING", (0, 0), (-1, -1), 0),
    ]))
    return t


def franja(valores, etiquetas, dentro, ancho_total, alto_mm=22):
    """Las barras de tendencia. Las de fuera del periodo van en tono claro: son
    contexto y no deben leerse como parte de lo que se informa."""
    maximo = max(valores) or 1
    ancho_col = ancho_total / max(len(valores), 1)
    centro = ParagraphStyle("c", parent=T, fontSize=6.6, leading=7.8, alignment=TA_CENTER)
    centro_m = ParagraphStyle("cm", parent=MUTE, alignment=TA_CENTER)
    cimas, barras, pies = [], [], []
    for v, et, es_del_periodo in zip(valores, etiquetas, dentro):
        alto = max(0.5, alto_mm * v / maximo)
        b = Table([[""]], colWidths=[ancho_col * 0.55], rowHeights=[alto * mm])
        b.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, -1),
             VERDE_MEDIO if es_del_periodo else VERDE_PALIDO),
            ("LEFTPADDING", (0, 0), (-1, -1), 0), ("RIGHTPADDING", (0, 0), (-1, -1), 0)]))
        caja = Table([[b]], colWidths=[ancho_col], rowHeights=[(alto_mm + 0.5) * mm])
        caja.setStyle(TableStyle([
            ("VALIGN", (0, 0), (-1, -1), "BOTTOM"), ("ALIGN", (0, 0), (-1, -1), "CENTER"),
            ("LEFTPADDING", (0, 0), (-1, -1), 0), ("RIGHTPADDING", (0, 0), (-1, -1), 0),
            ("TOPPADDING", (0, 0), (-1, -1), 0), ("BOTTOMPADDING", (0, 0), (-1, -1), 0)]))
        color = "#23241F" if es_del_periodo else "#9A968A"
        cimas.append(_p(f'<font color="{color}"><b>{v}</b></font>', centro))
        barras.append(caja)
        pies.append(_p(et, centro_m))
    t = Table([cimas, barras, pies], colWidths=[ancho_col] * len(valores))
    t.setStyle(TableStyle([
        ("VALIGN", (0, 1), (-1, 1), "BOTTOM"),
        ("LINEBELOW", (0, 1), (-1, 1), 0.5, LINEA),
        ("TOPPADDING", (0, 0), (-1, -1), 0.5), ("BOTTOMPADDING", (0, 0), (-1, -1), 0.5),
        ("LEFTPADDING", (0, 0), (-1, -1), 0), ("RIGHTPADDING", (0, 0), (-1, -1), 0),
    ]))
    t.hAlign = "LEFT"
    return t


# ---------------------------------------------------------------------------
# El PDF, armado desde la definición del área
# ---------------------------------------------------------------------------

def documento(spec):
    """El informe de una hoja. 'spec' lo arma cada área (ver informes_areas.py)."""
    desde, hasta = spec["desde"], spec["hasta"]
    p_desde, p_hasta = periodo_anterior(desde, hasta)
    dias = (datetime.date.fromisoformat(hasta)
            - datetime.date.fromisoformat(desde)).days + 1

    buf = io.BytesIO()
    doc = BaseDocTemplate(
        buf, pagesize=A4, topMargin=11 * mm, bottomMargin=11 * mm,
        leftMargin=MARGEN, rightMargin=MARGEN,
        title=f"{spec['titulo']} · {titulo_del_periodo(desde, hasta)}")
    marco = Frame(MARGEN, 11 * mm, ANCHO, A4[1] - 22 * mm, id="c",
                  leftPadding=0, rightPadding=0, topPadding=0, bottomPadding=0)

    def pie(canvas, d):
        canvas.saveState()
        canvas.setStrokeColor(LINEA)
        canvas.setLineWidth(0.5)
        canvas.line(MARGEN, 8.5 * mm, A4[0] - MARGEN, 8.5 * mm)
        canvas.setFont("Helvetica", 6.2)
        canvas.setFillColor(GRIS)
        canvas.drawString(MARGEN, 5.5 * mm, PIE_LIBRO)
        canvas.drawRightString(
            A4[0] - MARGEN, 5.5 * mm,
            f"Generated {datetime.datetime.now().strftime('%d/%m/%Y %H:%M')}")
        canvas.restoreState()

    doc.addPageTemplates([PageTemplate(id="u", frames=[marco], onPage=pie)])
    e = []

    # ---- Cabecera ----
    izq = []
    logo = _logo()
    if logo:
        img = Image(logo, width=30 * mm, height=9.5 * mm)
        img.hAlign = "LEFT"
        izq.append(img)
    izq += [Spacer(1, 2),
            _p(f'<font size="13" color="#2C4A3E"><b>{spec["titulo"]}</b></font>',
               ParagraphStyle("h", parent=T, fontSize=13, leading=15)),
            _p(spec.get("subtitulo") or "", MUTE)]
    der = [_p(f'<b>{titulo_del_periodo(desde, hasta)}</b>'
              + (f' · {dias} days' if dias > 1 else ""),
              ParagraphStyle("d1", parent=T, fontSize=8.5, leading=11, alignment=TA_RIGHT)),
           _p(f'compared with {fecha_larga(p_desde, False)} — {fecha_larga(p_hasta)}',
              ParagraphStyle("d2", parent=MUTE, alignment=TA_RIGHT))]
    cab = Table([[izq, der]], colWidths=[COL, COL + HUECO])
    cab.setStyle(TableStyle([
        ("VALIGN", (0, 0), (0, 0), "TOP"), ("VALIGN", (1, 0), (1, 0), "BOTTOM"),
        ("LEFTPADDING", (0, 0), (-1, -1), 0), ("RIGHTPADDING", (0, 0), (-1, -1), 0),
        ("TOPPADDING", (0, 0), (-1, -1), 0), ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ("LINEBELOW", (0, 0), (-1, 0), 1.2, BOSQUE),
    ]))
    e += [cab, Spacer(1, 5)]

    # ---- Cifras ----
    if spec.get("kpis"):
        e.append(banda_kpi(spec["kpis"], spec.get("neutrales", ())))
        linea = Table([[""]], colWidths=[ANCHO], rowHeights=[0.5])
        linea.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), LINEA)]))
        e += [Spacer(1, 1), linea, Spacer(1, 6)]

    # ---- Bloques, en dos columnas ----
    def render(b, ancho_col):
        piezas = seccion(b["titulo"], b.get("nota"), ancho=ancho_col)
        piezas.append(tablita(b["columnas"], b["filas"], b["anchos"],
                              total=b.get("total"), barra_en=b.get("barra_en"),
                              tope=b.get("tope")))
        if b.get("pie"):
            piezas += [Spacer(1, 1.5), _p(b["pie"], MUTE)]
        return piezas

    izq_b, der_b = [], []
    for b in spec.get("bloques", []):
        destino = der_b if b.get("lado") == "der" else izq_b
        if destino:
            destino.append(Spacer(1, 6))
        destino.extend(render(b, COL))
    if izq_b or der_b:
        cuerpo = Table([[izq_b, der_b]], colWidths=[COL, COL], hAlign="LEFT")
        cuerpo.setStyle(TableStyle([
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("LEFTPADDING", (0, 0), (0, 0), 0), ("RIGHTPADDING", (0, 0), (0, 0), HUECO),
            ("LEFTPADDING", (1, 0), (1, 0), 0), ("RIGHTPADDING", (1, 0), (1, 0), 0),
            ("TOPPADDING", (0, 0), (-1, -1), 0), ("BOTTOMPADDING", (0, 0), (-1, -1), 0),
        ]))
        e += [cuerpo, Spacer(1, 7)]

    # ---- Tendencia ----
    t = spec.get("tendencia")
    if t and t["valores"]:
        nota = (None if all(t["dentro"])
                else "In light tone, days before the period: shown as context only.")
        e += seccion(t["titulo"], nota, ancho=ANCHO)
        e.append(franja(t["valores"], t["etiquetas"], t["dentro"], ANCHO))
        e.append(Spacer(1, 7))

    # ---- Lectura ----
    if spec.get("lectura"):
        e += seccion(spec.get("titulo_lectura") or "Period summary", ancho=ANCHO)
        caja = Table([[_p(spec["lectura"],
                          ParagraphStyle("lec", parent=T, fontSize=7.6, leading=10.5))]],
                     colWidths=[ANCHO])
        caja.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, -1), ARENA),
            ("BOX", (0, 0), (-1, -1), 0.5, LINEA),
            ("LEFTPADDING", (0, 0), (-1, -1), 7), ("RIGHTPADDING", (0, 0), (-1, -1), 7),
            ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ]))
        e += [caja, Spacer(1, 3)]
    e.append(_p(spec.get("alcance") or ALCANCE, MUTE))

    doc.build(e)
    buf.seek(0)
    return buf


# ---------------------------------------------------------------------------
# El libro de Excel, armado desde la MISMA definición
# ---------------------------------------------------------------------------

_PROHIBIDO_EN_HOJA = re.compile(r"[\\/*?:\[\]]")


def _nombre_de_hoja(titulo, usados):
    """Excel no admite / \\ ? * : [ ] y corta en 31. Y no permite dos hojas iguales."""
    base = _PROHIBIDO_EN_HOJA.sub("-", titulo).strip()[:31] or "Sheet"
    nombre, i = base, 2
    while nombre.lower() in usados:
        sufijo = f" {i}"
        nombre = base[:31 - len(sufijo)] + sufijo
        i += 1
    usados.add(nombre.lower())
    return nombre


def libro(spec):
    """El mismo informe, en varias hojas y con los números como NÚMEROS.

    El PDF es para leer y presentar; esto es para filtrar y cruzar. Sale de la MISMA
    definición, así que lleva los mismos bloques, en el mismo orden y con las mismas
    columnas: una hoja por bloque, más el resumen, la tendencia, los datos y el glosario.
    """
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, Reference
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.table import Table as XTable, TableStyleInfo

    desde, hasta = spec["desde"], spec["hasta"]
    p_desde, p_hasta = periodo_anterior(desde, hasta)
    VERDE = "2C4A3E"
    wb = Workbook()
    usados = set()

    def hoja(titulo, encabezado, bajada=None):
        ws = wb.create_sheet(_nombre_de_hoja(titulo, usados))
        ws["A1"] = encabezado
        ws["A1"].font = Font(size=13, bold=True, color=VERDE)
        ws["A2"] = bajada or titulo_del_periodo(desde, hasta)
        ws["A2"].font = Font(size=9, italic=True, color="6E6A5C")
        return ws

    def volcar(ws, fila, columnas, filas, totales=None):
        for i, c in enumerate(columnas, start=1):
            celda = ws.cell(row=fila, column=i, value=c[1])
            celda.font = Font(bold=True, color="FFFFFF")
            celda.fill = PatternFill("solid", fgColor=VERDE)
        ws.freeze_panes = ws.cell(row=fila + 1, column=1)
        r = fila
        for f in filas:
            r += 1
            for i, c in enumerate(columnas, start=1):
                celda = ws.cell(row=r, column=i, value=f.get(c[0]))
                if len(c) > 3 and c[3]:
                    celda.number_format = c[3]
        if totales:
            r += 1
            borde = Border(top=Side(style="thin", color=VERDE))
            for i, c in enumerate(columnas, start=1):
                celda = ws.cell(row=r, column=i, value=totales.get(c[0]))
                celda.font = Font(bold=True)
                celda.border = borde
                if len(c) > 3 and c[3]:
                    celda.number_format = c[3]
        for i, c in enumerate(columnas, start=1):
            largos = [len(str(c[1]))] + [len(str(f.get(c[0]) or "")) for f in filas]
            ws.column_dimensions[get_column_letter(i)].width = min(max(largos) + 3, 45)
        return r

    # ---- Summary ----
    ws = hoja("Summary", spec["titulo"],
              f"{titulo_del_periodo(desde, hasta)} · compared with "
              f"{fecha_larga(p_desde, False)} — {fecha_larga(p_hasta)}")
    cols = [("kpi", "Indicator", "i", None), ("actual", "Period", "d", "#,##0"),
            ("previo", "Previous", "d", "#,##0"), ("var", "Change", "d", "0.0%")]
    filas = []
    for etiqueta, valor, previo, sufijo, _pie in spec.get("kpis", []):
        filas.append({"kpi": f"{etiqueta}{f' ({sufijo})' if sufijo else ''}",
                      "actual": valor, "previo": previo,
                      "var": None if not previo else (valor - previo) / previo})
    volcar(ws, 4, cols, filas)
    if spec.get("lectura"):
        fila = 6 + len(filas)
        ws.cell(row=fila, column=1, value=spec.get("titulo_lectura") or "Period summary"
                ).font = Font(bold=True, color=VERDE)
        limpio = re.sub(r"<[^>]+>", "", spec["lectura"])
        celda = ws.cell(row=fila + 1, column=1, value=limpio)
        celda.alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=fila + 1, start_column=1, end_row=fila + 4, end_column=4)

    # ---- Una hoja por bloque, en el MISMO orden del PDF ----
    for b in spec.get("bloques", []):
        ws = hoja(b["titulo"], b["titulo"], b.get("nota"))
        columnas = [c for c in b["columnas"] if c[0] != b.get("barra_en")]
        volcar(ws, 4, columnas, b["filas"], b.get("total"))
        if b.get("grafico") and b["filas"]:
            ch = BarChart()
            ch.type, ch.legend = "bar", None
            ch.title = b["titulo"]
            ch.height, ch.width = 8, 14
            col_valor = next(i for i, c in enumerate(columnas, start=1)
                             if c[0] == b["grafico"])
            ch.add_data(Reference(ws, min_col=col_valor, min_row=4,
                                  max_row=4 + len(b["filas"])), titles_from_data=True)
            ch.set_categories(Reference(ws, min_col=1, min_row=5,
                                        max_row=4 + len(b["filas"])))
            ws.add_chart(ch, f"{get_column_letter(len(columnas) + 2)}4")

    # ---- Tendencia ----
    t = spec.get("tendencia")
    if t and t["valores"]:
        ws = hoja("Trend", t["titulo"])
        volcar(ws, 4,
               [("periodo", "Period", "i", None), ("valor", "Value", "d", "#,##0"),
                ("dentro", "In report period", "i", None)],
               [{"periodo": et, "valor": v, "dentro": "yes" if d else "context"}
                for v, et, d in zip(t["valores"], t["etiquetas"], t["dentro"])])

    # ---- Datos ----
    d = spec.get("datos")
    if d and d["filas"]:
        ws = hoja("Data", d.get("titulo") or "One row per record",
                  "This is the sheet for pivot tables: it has a filter and table format.")
        fin = volcar(ws, 4, d["columnas"], d["filas"])
        tab = XTable(displayName="Records",
                     ref=f"A4:{get_column_letter(len(d['columnas']))}{fin}")
        tab.tableStyleInfo = TableStyleInfo(name="TableStyleLight11", showRowStripes=True)
        ws.add_table(tab)

    # ---- Glosario ----
    definiciones = list(spec.get("definiciones") or [])
    definiciones += [
        ("Previous period", "A full calendar month is compared with the previous "
                            "calendar month; any other range, with the same number of "
                            "days immediately before."),
        (SIN_ASIGNAR, "The system does not have that value. It is not a zero: it is "
                      "something missing."),
    ]
    ws = hoja("Glossary", "What each figure means",
              "So that two people do not read the same number differently.")
    ws["A4"], ws["B4"] = "Term", "Definition"
    for c in ("A4", "B4"):
        ws[c].font = Font(bold=True, color="FFFFFF")
        ws[c].fill = PatternFill("solid", fgColor=VERDE)
    for i, (termino, texto) in enumerate(definiciones, start=5):
        ws.cell(row=i, column=1, value=termino).font = Font(bold=True)
        ws.cell(row=i, column=2, value=texto).alignment = Alignment(wrap_text=True,
                                                                   vertical="top")
        ws.row_dimensions[i].height = 26
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 95

    del wb["Sheet"]
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
