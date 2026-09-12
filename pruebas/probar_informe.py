"""Los informes: una hoja, el mismo formato en todas las areas, y en ingles.

Lo que se descargaba antes era una tabla plana de tres columnas, la misma para todo, sin
totales ni comparacion, y con secciones enteras fuera. Ahora cada area declara su informe
UNA vez y de ahi salen el PDF y el libro de Excel.

Lo que estas pruebas fijan, que es lo facil de romper sin darse cuenta:

  1. UNA pagina, en TODAS las areas. Basta con que un mes traiga un guia mas para que se
     desborde a la segunda hoja.
  2. El MISMO formato por dia y por mes.
  3. El PDF y el Excel salen de la MISMA definicion. Si divergieran, nadie sabria cual
     creer.
  4. Todo en INGLES: se mandan afuera.
"""
import datetime
import os
import re
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import comun
from comun import cargar, reserva

import informe as inf
import informes_areas as areas

HOY = datetime.date.today()


def ddmmyy(d):
    return f"{d.day:02d}-{d.month:02d}-{str(d.year)[2:]}"


def _sembrar(conn, dias=40, cuantas=30, cortesias=4):
    """Un periodo con de todo, para que los informes tengan que resumir."""
    import random
    random.seed(5)
    tours = ["PNC", "SIRENA", "SNORKEL", "MANGLAR", "NW", "BALLENAS", "CABALGATA"]
    inicio = HOY - datetime.timedelta(days=dias)
    lote = []
    for i in range(cuantas):
        llega = inicio + datetime.timedelta(days=random.randint(0, dias - 4))
        noches = random.choice([2, 3, 4])
        lote.append(reserva(
            f"80{i:04d}", f"{(i % 20) + 1:02d}", ddmmyy(llega),
            ddmmyy(llega + datetime.timedelta(days=noches)),
            estado="EN CASA", adl=random.choice([1, 2, 2, 3]),
            punto_entrada=random.choice(["Sierpe", "Drake"]),
            punto_salida=random.choice(["Sierpe", "Drake"]),
            # Ni la cortesia ni las amenidades se ponen a mano: se escriben en el texto
            # y las detecta el importador, que es como llegan de Opera.
            texto_completo=" ".join([
                "PAQUETE 3D/2N+ FULLBOARD CPL EN HOSPEDAJE" if i < cortesias
                else "PENSION COMPLETA",
                random.choice(["", "", "Gluten free diet", "no seafood",
                               "Honeymoon - luna de miel", "Vegetarian"]),
            ]).strip(),
            operacion=[{"dia": (llega + datetime.timedelta(days=d)).day,
                        "tour": random.choice(tours)}
                       for d in range(1, min(noches, 3))]))
    cargar(lote)
    botes = [r["nombre"] for r in conn.execute("SELECT nombre FROM bote")][:4]
    guias = [r["nombre"] for r in conn.execute("SELECT nombre FROM guia")][:6]
    for f in conn.execute("SELECT id FROM tour_asignado").fetchall():
        conn.execute("UPDATE tour_asignado SET bote_nombre=?, guia_nombre=? WHERE id=?",
                     (random.choice(botes), random.choice(guias), f["id"]))
    servicio = conn.execute("SELECT codigo FROM spa_servicio LIMIT 1").fetchone()
    terapeuta = conn.execute("SELECT nombre FROM spa_terapeuta LIMIT 1").fetchone()
    confs = [r["conf_no"] for r in conn.execute("SELECT conf_no FROM reserva")]
    for i in range(25):
        conn.execute(
            """INSERT INTO spa_cita (conf_no, fecha, servicio_codigo, minutos,
                                     terapeuta, estado, origen)
               VALUES (?,?,?,?,?,?,'HUESPED')""",
            (random.choice(confs),
             (inicio + datetime.timedelta(days=random.randint(0, dias))).isoformat(),
             servicio["codigo"], random.choice([50, 60, 80]),
             terapeuta["nombre"] if terapeuta else None,
             random.choice(["HECHA", "CONFIRMADA", "SOLICITADA", "CANCELADA"])))
    conn.commit()


def _paginas(buf):
    import pdfplumber
    buf.seek(0)
    with pdfplumber.open(buf) as pdf:
        return len(pdf.pages), "\n".join((p.extract_text() or "") for p in pdf.pages)


def _mes_pasado():
    fin = HOY.replace(day=1) - datetime.timedelta(days=1)
    return fin.replace(day=1).isoformat(), fin.isoformat()


# ---------------------------------------------------------------------------
# Que quepan, TODAS
# ---------------------------------------------------------------------------

def todas_las_areas_caben_en_una_hoja(c):
    """Es el requisito, y vale para las ocho: un informe de dos hojas ya no es un
    informe de una hoja."""
    conn = comun.base_limpia()
    _sembrar(conn)
    ini, fin = _mes_pasado()
    for area in areas.AREAS:
        n, _ = _paginas(areas.construir(area, conn, ini, fin, "pdf"))
        c.igual(n, 1, f"«{area}» de un mes ocupa UNA hoja")
    conn.close()


def todas_las_areas_caben_tambien_por_dia(c):
    conn = comun.base_limpia()
    _sembrar(conn)
    for area in areas.AREAS:
        n, _ = _paginas(areas.construir(area, conn, HOY.isoformat(), HOY.isoformat(), "pdf"))
        c.igual(n, 1, f"«{area}» de un dia ocupa UNA hoja")
    conn.close()


def un_periodo_sin_nada_no_revienta(c):
    """Un dia sin movimiento tiene que dar un informe que lo diga, no un error ni una
    hoja rota: se va a pedir en temporada baja."""
    conn = comun.base_limpia()
    lejos = (HOY + datetime.timedelta(days=300)).isoformat()
    for area in areas.AREAS:
        n, texto = _paginas(areas.construir(area, conn, lejos, lejos, "pdf"))
        c.igual(n, 1, f"«{area}» sin datos sigue siendo una hoja")
        c.cierto(len(texto) > 50, f"«{area}» sin datos sigue diciendo algo")
    conn.close()


# ---------------------------------------------------------------------------
# Que sea el MISMO formato
# ---------------------------------------------------------------------------

def el_formato_es_el_mismo_por_dia_y_por_mes(c):
    conn = comun.base_limpia()
    _sembrar(conn)
    ini, fin = _mes_pasado()
    _, mes = _paginas(areas.construir("analitica", conn, ini, fin, "pdf"))
    _, dia = _paginas(areas.construir("analitica", conn, HOY.isoformat(),
                                      HOY.isoformat(), "pdf"))
    for bloque in ("TOURS BY TYPE", "GUIDES", "BOATS", "EMBARKATION POINTS", "SPA",
                   "PERIOD SUMMARY", "PAX-NIGHTS", "COMPLIMENTARY"):
        c.cierto(bloque in mes, f"«{bloque}» esta en el informe del mes")
        c.cierto(bloque in dia, f"«{bloque}» esta tambien en el del dia")
    conn.close()


def todas_las_areas_tienen_la_misma_estructura(c):
    """Es lo que se pidio: el informe de Restaurantes se lee igual que el de Analitica."""
    conn = comun.base_limpia()
    _sembrar(conn)
    ini, fin = _mes_pasado()
    for area in areas.AREAS:
        spec = areas.AREAS[area](conn, ini, fin)
        c.cierto(spec.get("titulo"), f"«{area}» tiene titulo")
        c.cierto(len(spec.get("kpis") or []) >= 4,
                 f"«{area}» tiene su banda de cifras")
        c.cierto(len(spec.get("bloques") or []) >= 2,
                 f"«{area}» tiene al menos dos bloques de tabla")
        c.cierto(spec.get("lectura"), f"«{area}» tiene su parrafo de lectura")
        c.cierto(spec.get("datos"), f"«{area}» tiene su tabla plana para el Excel")
        for b in spec["bloques"]:
            c.igual(len(b["anchos"]), len(b["columnas"]),
                    f"«{area} / {b['titulo']}»: un ancho por columna")
    conn.close()


def el_pdf_y_el_excel_salen_de_la_misma_definicion(c):
    """Si salieran de dos codigos distintos, uno se quedaria viejo y nadie sabria cual
    creer. El libro lleva una hoja por bloque, con el mismo nombre y en el mismo orden."""
    import openpyxl
    conn = comun.base_limpia()
    _sembrar(conn)
    ini, fin = _mes_pasado()
    spec = areas.AREAS["analitica"](conn, ini, fin)
    wb = openpyxl.load_workbook(areas.construir("analitica", conn, ini, fin, "xlsx"))
    hojas = wb.sheetnames
    c.igual(hojas[0], "Summary", "el libro abre con el resumen")
    for i, b in enumerate(spec["bloques"]):
        c.igual(hojas[i + 1], b["titulo"][:31],
                f"la hoja {i + 2} es «{b['titulo']}», igual que en el PDF")
    for esperada in ("Trend", "Data", "Glossary"):
        c.cierto(esperada in hojas, f"y el libro trae «{esperada}»")
    ws = wb["Data"]
    c.cierto("Records" in list(ws.tables),
             "la hoja de datos es una TABLA de Excel, para poder pivotear")
    c.cierto(ws.freeze_panes is not None, "con el encabezado fijo")
    conn.close()


def el_excel_guarda_numeros_no_texto(c):
    import openpyxl
    conn = comun.base_limpia()
    _sembrar(conn)
    ini, fin = _mes_pasado()
    wb = openpyxl.load_workbook(areas.construir("analitica", conn, ini, fin, "xlsx"))
    ws = wb["Summary"]
    c.cierto(any(isinstance(ws.cell(row=r, column=2).value, (int, float))
                 for r in range(5, 12)), "los valores son numeros")
    c.igual(ws.cell(row=5, column=4).number_format, "0.0%",
            "y la variacion lleva formato de porcentaje")
    conn.close()


# ---------------------------------------------------------------------------
# En ingles
# ---------------------------------------------------------------------------

def los_informes_salen_en_ingles(c):
    """Se mandan a inversionistas y socios que no leen espanol. La pantalla sigue en
    espanol; lo que cambia es lo que se imprime."""
    conn = comun.base_limpia()
    _sembrar(conn)
    ini, fin = _mes_pasado()
    # Palabras que delatan que quedo texto en espanol en la parte que se IMPRIME. No se
    # buscan nombres propios ni datos: Sierpe, Drake o los tours se quedan como estan.
    delatoras = re.compile(
        r"\b(Informe|Reporte|Periodo|Resumen|Guías|Botes|Cortesías|Comparado|"
        r"Generado|Página|Fecha|Huésped|Salidas|Tratamiento|Alcance)\b")
    for area in areas.AREAS:
        _, texto = _paginas(areas.construir(area, conn, ini, fin, "pdf"))
        encontradas = sorted(set(delatoras.findall(texto)))
        c.igual(encontradas, [], f"«{area}» no imprime palabras en espanol")
    conn.close()


def los_meses_y_los_estados_van_en_ingles(c):
    c.igual(inf.titulo_del_periodo("2026-03-01", "2026-03-31"), "March 2026",
            "un mes completo se dice por su nombre, en ingles")
    c.igual(inf.titulo_del_periodo("2026-03-10", "2026-03-10"), "March 10, 2026",
            "un dia, con su fecha")
    c.igual(areas._ingles_estado("HECHA"), "completed", "los estados se traducen")
    c.igual(areas._ingles_estado("SOLICITADA"), "requested", "todos")


# ---------------------------------------------------------------------------
# Lo que cambia entre un dia y un mes
# ---------------------------------------------------------------------------

def un_mes_se_compara_con_el_mes_anterior(c):
    """Comparar un mes contra «los 30 dias de antes» corre el corte y mezcla dos meses."""
    c.igual(inf.periodo_anterior("2026-03-01", "2026-03-31"),
            ("2026-02-01", "2026-02-28"), "marzo se compara con febrero entero")
    c.igual(inf.periodo_anterior("2026-01-01", "2026-01-31"),
            ("2025-12-01", "2025-12-31"), "enero, con diciembre del ano anterior")
    c.igual(inf.periodo_anterior("2026-03-05", "2026-03-11"),
            ("2026-02-26", "2026-03-04"), "una semana, con la semana de antes")
    c.igual(inf.periodo_anterior("2026-03-10", "2026-03-10"),
            ("2026-03-09", "2026-03-09"), "un dia, con el dia anterior")


def la_tendencia_se_adapta_al_largo_del_periodo(c):
    """«Por semana» no significa nada en un informe de un dia."""
    conn = comun.base_limpia()
    _sembrar(conn)
    _, _, _, titulo_mes = inf.serie_de_tendencia(
        conn, (HOY - datetime.timedelta(days=29)).isoformat(), HOY.isoformat())
    c.cierto("week" in titulo_mes, "un mes se grafica por semana")

    valores, _, dentro, titulo_dia = inf.serie_de_tendencia(
        conn, HOY.isoformat(), HOY.isoformat())
    c.cierto("day" in titulo_dia, "un dia se grafica por dia")
    c.cierto(len(valores) >= inf.MINIMO_BARRAS,
             "y se completa hacia atras para tener con que comparar")
    c.igual(dentro.count(True), 1, "solo la barra del dia pedido es del periodo")
    c.igual(dentro[-1], True, "y es la ultima")
    conn.close()


# ---------------------------------------------------------------------------
# Que los numeros sean los de la pantalla
# ---------------------------------------------------------------------------

def ninguna_tendencia_sale_entera_en_cero(c):
    """El fallo que costo mas de encontrar: las consultas que agrupan por la fecha de la
    reserva devuelven la clave en YY-MM-DD y la serie la busca en ISO. Sin convertirla no
    encontraba NADA y el grafico salia entero en cero, sin dar ningun error — que es la
    peor forma de fallar, porque un cero se lee como «no hubo nadie»."""
    conn = comun.base_limpia()
    _sembrar(conn)
    ini, fin = _mes_pasado()
    for area in areas.AREAS:
        spec = areas.AREAS[area](conn, ini, fin)
        t = spec.get("tendencia")
        if not t or not t["valores"]:
            continue
        c.cierto(sum(t["valores"]) > 0,
                 f"la tendencia de «{area}» tiene datos, no puros ceros")
    conn.close()


def las_etiquetas_caben_en_su_casilla(c):
    """Con siete casillas cada una tiene 26 mm. Una etiqueta larga se parte en dos
    renglones y se monta encima del numero."""
    conn = comun.base_limpia()
    _sembrar(conn)
    ini, fin = _mes_pasado()
    for area in areas.AREAS:
        spec = areas.AREAS[area](conn, ini, fin)
        ancho = inf.ANCHO / len(spec["kpis"]) / inf.mm
        for etiqueta, *_ in spec["kpis"]:
            # ~1.55 mm por caracter a 6.2 pt en mayusculas, medido sobre la hoja.
            c.cierto(len(etiqueta) * 1.55 <= ancho + 1,
                     f"«{area}»: la etiqueta «{etiqueta}» cabe en {ancho:.0f} mm")
        for b in spec["bloques"]:
            for i, col in enumerate(b["columnas"]):
                if col[0] == b.get("barra_en"):
                    continue
                mm_col = b["anchos"][i] / inf.mm
                c.cierto(len(col[1]) * 1.35 <= mm_col + 2,
                         f"«{area} / {b['titulo']}»: «{col[1]}» cabe en {mm_col:.0f} mm")
    conn.close()


def las_cortesias_se_cuentan_no_se_valoran(c):
    conn = comun.base_limpia()
    _sembrar(conn, cortesias=4)
    m = inf.metricas(conn, (HOY - datetime.timedelta(days=40)).isoformat(),
                     HOY.isoformat())
    c.igual(m["cortesias"], 4, "cuenta las reservas marcadas como cortesia")
    c.cierto(m["cortesias_pax"] >= 4, "y su gente")
    conn.close()


def la_ocupacion_usa_el_mismo_total_que_la_pantalla(c):
    """Si el informe contara las habitaciones de otra forma, diria un porcentaje distinto
    al de la pantalla de ocupacion y nadie sabria cual creer."""
    conn = comun.base_limpia()
    llega = HOY - datetime.timedelta(days=1)
    cargar([reserva(f"90{i}", f"{i:02d}", ddmmyy(llega),
                    ddmmyy(HOY + datetime.timedelta(days=2)), estado="EN CASA", adl=2)
            for i in range(1, 11)])
    m = inf.metricas(conn, HOY.isoformat(), HOY.isoformat(), habitaciones=20)
    c.igual(m["ocupacion"], 50, "10 habitaciones de 20 es 50%")
    c.igual(m["pax_noche"], 20, "y 20 pax durmiendo")
    conn.close()


def el_spa_si_entra_en_el_informe(c):
    """El reporte viejo dejaba el spa fuera aunque la pantalla lo mostraba."""
    conn = comun.base_limpia()
    _sembrar(conn)
    ini, fin = _mes_pasado()
    _, texto = _paginas(areas.construir("analitica", conn, ini, fin, "pdf"))
    c.cierto("SPA" in texto, "el spa aparece en la hoja")
    conn.close()


def un_area_inventada_se_rechaza(c):
    conn = comun.base_limpia()
    try:
        areas.construir("no-existe", conn, HOY.isoformat(), HOY.isoformat())
        c.cierto(False, "un area desconocida tiene que rebotar")
    except ValueError:
        c.ok += 1
    conn.close()


PRUEBAS = [
    todas_las_areas_caben_en_una_hoja,
    todas_las_areas_caben_tambien_por_dia,
    un_periodo_sin_nada_no_revienta,
    el_formato_es_el_mismo_por_dia_y_por_mes,
    todas_las_areas_tienen_la_misma_estructura,
    el_pdf_y_el_excel_salen_de_la_misma_definicion,
    el_excel_guarda_numeros_no_texto,
    los_informes_salen_en_ingles,
    los_meses_y_los_estados_van_en_ingles,
    un_mes_se_compara_con_el_mes_anterior,
    la_tendencia_se_adapta_al_largo_del_periodo,
    ninguna_tendencia_sale_entera_en_cero,
    las_etiquetas_caben_en_su_casilla,
    las_cortesias_se_cuentan_no_se_valoran,
    la_ocupacion_usa_el_mismo_total_que_la_pantalla,
    el_spa_si_entra_en_el_informe,
    un_area_inventada_se_rechaza,
]


if __name__ == "__main__":
    sys.exit(comun.correr("informes", PRUEBAS))
